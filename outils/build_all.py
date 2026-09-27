#!/usr/bin/env python3
"""
Génère tous les fichiers du cours dans build/ :

  build/slides/         PowerPoint modifiables (.pptx) + PDF, une présentation par séance
  build/enseignant/     Word (.docx) + PDF : notes, scripts de démo, grilles, jalons, plan
  build/etudiants/      Word (.docx) + PDF : énoncés, aide-mémoire, cahier des charges
  build/code/           ZIP : starter et corrigé de chaque TP et atelier, démos

Usage (depuis la racine du dépôt) :
  python3 outils/build_all.py            tout
  python3 outils/build_all.py slides     seulement les slides
  python3 outils/build_all.py docs       seulement les documents
  python3 outils/build_all.py code       seulement les ZIP

Outils nécessaires : node + pptxgenjs (npm install -g pptxgenjs), pandoc, LibreOffice (soffice).
"""
import os
import re
import shlex
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / "build"
AUTHOR = "Amani Selmi, PhD-Engineer"
COURSE = "Développement Web 1"

# Dossier -> (public, nom court) ; le public décide si le document va chez l'enseignante ou chez les étudiants
TEACHER = "enseignant"
STUDENT = "etudiants"


def run(cmd, **kwargs):
    subprocess.run(cmd, check=True, **kwargs)


def soffice_to_pdf(files, outdir):
    """Convertit une liste de fichiers Office en PDF (LibreOffice, par lots)."""
    files = [str(f) for f in files]
    for i in range(0, len(files), 20):
        # SOFFICE_CMD permet de remplacer la commande (par défaut : soffice)
        soffice = shlex.split(os.environ.get("SOFFICE_CMD", "soffice"))
        run([*soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), *files[i:i + 20]],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=900)


# ------------------------------------------------------------------ slides
def build_slides():
    out = BUILD / "slides"
    out.mkdir(parents=True, exist_ok=True)
    decks = []
    for md in sorted((ROOT / "semaines").glob("S*/slides.md")):
        week = md.parent.name                       # "S03-javascript"
        target = out / f"{week}.pptx"
        run(["node", str(ROOT / "outils" / "build_slides.js"), str(md), str(target)])
        decks.append(target)
    soffice_to_pdf(decks, out)
    print(f"{len(decks)} présentations → {out}")


# ------------------------------------------------------------------ documents
def document_list():
    """(fichier source, public, nom de sortie)"""
    docs = [
        (ROOT / "COURSE_BIBLE.md", TEACHER, "00-course-bible"),
        (ROOT / "plan" / "plan-semestre.md", TEACHER, "00-plan-du-cours"),
        (ROOT / "projet" / "soutenance.md", TEACHER, "projet-guide-soutenance"),
        (ROOT / "projet" / "specification.md", STUDENT, "projet-cahier-des-charges"),
        (ROOT / "projet" / "proposition-modele.md", STUDENT, "projet-proposition-modele"),
        (ROOT / "projet" / "grille.md", STUDENT, "projet-grille-de-notation"),
    ]
    for fiche in sorted((ROOT / "fiches").glob("*.md")):
        docs.append((fiche, STUDENT, f"fiche-{fiche.stem}"))

    teacher_names = {"notes-enseignant.md": "notes", "live-coding.md": "live-coding", "grille.md": "grille",
                     "jalon1.md": "jalon1", "jalon2.md": "jalon2", "quiz-diagnostic.md": "quiz-diagnostic",
                     "planning-modele.md": "planning-soutenances"}
    student_names = {"enonce.md": "enonce", "revue-de-code.md": "revue-de-code",
                     "preparation-soutenance.md": "preparation-soutenance"}

    for week in sorted((ROOT / "semaines").glob("S*")):
        code = week.name.split("-")[0]              # "S03"
        for md in sorted(week.rglob("*.md")):
            rel = md.relative_to(week)
            if "starter" in rel.parts or "solution" in rel.parts or md.name in ("README.md", "slides.md"):
                continue
            part = rel.parts[0] if len(rel.parts) > 1 else ""   # "tp2", "atelier" ou ""
            prefix = f"{code}-{part}-" if part else f"{code}-"
            if md.name in teacher_names:
                docs.append((md, TEACHER, prefix + teacher_names[md.name]))
            elif md.name in student_names:
                docs.append((md, STUDENT, prefix + student_names[md.name]))
    return docs


def make_reference_docx(path):
    """Styles Word du cours : Calibri, titres bleus, code en Courier New."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    with open(path, "wb") as f:
        f.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                               check=True, capture_output=True).stdout)
    doc = Document(path)
    blue = RGBColor(0x1D, 0x4E, 0xD8)
    ink = RGBColor(0x1E, 0x29, 0x3B)
    styles = doc.styles
    for name, size, color, bold in [("Normal", 11, ink, False), ("Body Text", 11, ink, False),
                                    ("First Paragraph", 11, ink, False), ("Compact", 11, ink, False),
                                    ("Title", 24, blue, True), ("Subtitle", 13, ink, False), ("Author", 12, ink, False),
                                    ("Heading 1", 16, blue, True), ("Heading 2", 13, blue, True), ("Heading 3", 12, ink, True)]:
        try:
            st = styles[name]
        except KeyError:
            continue
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = bold
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if fonts is not None:
            for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + attr, None)
    for name in ("Title", "Author", "Subtitle"):
        try:
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        except KeyError:
            pass
    for name in ("Source Code", "Verbatim Char"):
        try:
            styles[name].font.name = "Courier New"
            styles[name].font.size = Pt(9.5)
        except KeyError:
            pass
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(60)
        section.top_margin = section.bottom_margin = Pt(55)
    doc.save(path)


def style_tables(path):
    """Tableaux lisibles : bordures fines, en-tête bleu, largeurs de colonnes selon le contenu."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, RGBColor

    doc = Document(path)
    usable = 6.6  # pouces
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "CBD5E1")
            borders.append(el)
        for old in tbl_pr.findall(qn("w:tblBorders")):
            tbl_pr.remove(old)
        tbl_pr.append(borders)

        rows = table.rows
        cols = len(table.columns)
        lengths = [max(4, min(45, max(len(r.cells[c].text) for r in rows))) for c in range(cols)]
        total = sum(lengths)
        widths = [max(0.6, usable * l / total) for l in lengths]
        scale = usable / sum(widths)
        widths = [w * scale for w in widths]
        grid = table._tbl.tblGrid
        for c, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[c] * 1440)))
        for r_index, row in enumerate(rows):
            for c, cell in enumerate(row.cells):
                cell.width = Inches(widths[c])
                if r_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:val"), "clear")
                    shading.set(qn("w:color"), "auto")
                    shading.set(qn("w:fill"), "1D4ED8")
                    cell._tc.get_or_add_tcPr().append(shading)
                    for paragraph in cell.paragraphs:
                        for run_ in paragraph.runs:
                            run_.font.bold = True
                            run_.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.save(path)


def build_docs():
    ref = BUILD / "reference.docx"
    BUILD.mkdir(exist_ok=True)
    make_reference_docx(ref)
    produced = {TEACHER: [], STUDENT: []}
    for src, audience, name in document_list():
        out_dir = BUILD / audience
        out_dir.mkdir(parents=True, exist_ok=True)
        target = out_dir / f"{name}.docx"
        text = src.read_text(encoding="utf-8")
        # Les liens vers d'autres fichiers du dépôt n'ont pas de sens dans un Word : on garde le texte
        text = re.sub(r"\[([^\]]+)\]\((?!https?://)[^)]+\)", r"\1", text)
        run(["pandoc", "-f", "gfm", "-t", "docx", "--shift-heading-level-by=-1",
             "--reference-doc", str(ref), "-M", f"author={AUTHOR}", "-M", f"subtitle={COURSE}",
             "-o", str(target)], input=text.encode("utf-8"))
        style_tables(target)
        produced[audience].append(target)
    for audience, files in produced.items():
        soffice_to_pdf(files, BUILD / audience)
        print(f"{len(files)} documents → {BUILD / audience}")
    ref.unlink()


# ------------------------------------------------------------------ code
def zip_dir(src, target, extra=None):
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as z:
        for path in sorted(src.rglob("*")):
            if path.is_file() and path.name != "config.php" and ".git" not in path.parts:
                z.write(path, Path(src.name) / path.relative_to(src))
        for folder in extra or []:
            for path in sorted(folder.rglob("*")):
                if path.is_file():
                    z.write(path, Path(src.name) / folder.name / path.relative_to(folder))


def build_code():
    out = BUILD / "code"
    out.mkdir(parents=True, exist_ok=True)
    count = 0
    for week in sorted((ROOT / "semaines").glob("S*")):
        code = week.name.split("-")[0]
        for kind in ("starter", "solution"):
            for folder in sorted(week.glob(f"*/{kind}")):
                extra = [folder.parent / "maquette"] if kind == "starter" and (folder.parent / "maquette").is_dir() else []
                zip_dir(folder, out / f"{code}-{folder.parent.name}-{kind}.zip", extra)
                count += 1
        demo = week / "live-coding" / "demo"
        if demo.is_dir():
            zip_dir(demo, out / f"{code}-demo.zip")
            count += 1
    print(f"{count} archives → {out}")


if __name__ == "__main__":
    what = sys.argv[1] if len(sys.argv) > 1 else "all"
    if what in ("all", "slides"):
        build_slides()
    if what in ("all", "docs"):
        build_docs()
    if what in ("all", "code"):
        build_code()
